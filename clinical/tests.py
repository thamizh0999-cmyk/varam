from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from .models import (
    CentreProfile,
    FacilityMembership,
    IPPatient,
    DailyCensus,
    DailyDiseaseCount,
    DiseaseMaster,
    IPDiseaseRecord,
    IPNCDRecord,
    ANCPatient,
    ANCPatient,
    ANCVisit,
    DailyNCDCount
)


class FacilityIsolationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.facility_a = CentreProfile.objects.create(
            centre_name="Test Facility A",
            district="Test District A",
            facility_type="GH",
            is_active=True,
        )
        cls.facility_b = CentreProfile.objects.create(
            centre_name="Test Facility B",
            district="Test District B",
            facility_type="GH",
            is_active=True,
        )

        cls.user_a = User.objects.create_user(
            username="facility_user_a",
            password="test-password-a",
        )
        cls.user_b = User.objects.create_user(
            username="facility_user_b",
            password="test-password-b",
        )

        FacilityMembership.objects.create(
            user=cls.user_a,
            facility=cls.facility_a,
            role="FACILITY_ADMIN",
            is_active=True,
        )
        FacilityMembership.objects.create(
            user=cls.user_b,
            facility=cls.facility_b,
            role="FACILITY_ADMIN",
            is_active=True,
        )

        cls.patient_a = IPPatient.objects.create(
            facility=cls.facility_a,
            name="Test Patient A",
            sex="M",
            is_child=False,
            admission_date=date(2026, 10, 10),
            admission_session="MORNING",
        )
        cls.patient_b = IPPatient.objects.create(
            facility=cls.facility_b,
            name="Test Patient B",
            sex="F",
            is_child=False,
            admission_date=date(2026, 10, 10),
            admission_session="MORNING",
        )

        cls.census_a = DailyCensus.objects.create(
            facility=cls.facility_a,
            date=date(2026, 10, 10),
        )
        cls.census_b = DailyCensus.objects.create(
            facility=cls.facility_b,
            date=date(2026, 10, 10),
        )
        cls.anc_patient_a = ANCPatient.objects.create(
            facility=cls.facility_a,
            name="ANC Patient Facility A",
            age=24,
            lmp=date(2026, 1, 1),
            edd=date(2026, 10, 8),
        )
        cls.anc_patient_b = ANCPatient.objects.create(
            facility=cls.facility_b,
            name="ANC Patient Facility B",
            age=26,
            lmp=date(2026, 1, 5),
            edd=date(2026, 10, 12),
        )

    def test_ip_register_shows_only_assigned_facility_patients(self):
        self.client.force_login(self.user_a)

        response = self.client.get(reverse("ip_register"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Test Patient A")
        self.assertNotContains(response, "Test Patient B")

    def test_cannot_edit_another_facility_patient(self):
        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("ip_edit", args=[self.patient_b.pk])
        )

        self.assertEqual(response.status_code, 404)

    def test_cannot_delete_another_facility_patient(self):
        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse("ip_delete", args=[self.patient_b.pk])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(
            IPPatient.objects.filter(pk=self.patient_b.pk).exists()
        )

    def test_daily_census_displays_only_assigned_facility(self):
        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("daily_census"),
            {"date": "2026-10-10"},
        )

        self.assertEqual(response.status_code, 200)

        context = response.context
        self.assertEqual(context["facility"].pk, self.facility_a.pk)

    
    def test_ip_disease_report_excludes_other_facility_records(self):
        disease_a = DiseaseMaster.objects.create(
            name="Facility A Test Disease",
            is_active=True,
        )
        disease_b = DiseaseMaster.objects.create(
            name="Facility B Test Disease",
            is_active=True,
        )

        IPDiseaseRecord.objects.create(
            patient=self.patient_a,
            disease=disease_a,
            diagnosis_date=date(2026, 10, 10),
        )
        IPDiseaseRecord.objects.create(
            patient=self.patient_b,
            disease=disease_b,
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("ip_disease_report"),
            {
                "start_date": "2026-10-01",
                "end_date": "2026-10-31",
            },
        )

        self.assertEqual(response.status_code, 200)

        rows = response.context["report_rows"]
        disease_names = [
            row["disease"].name for row in rows
        ]

        self.assertIn("Facility A Test Disease", disease_names)
        self.assertNotIn("Facility B Test Disease", disease_names)

        row_a = next(
            row for row in rows
            if row["disease"].name == "Facility A Test Disease"
        )
        self.assertEqual(row_a["total"], 1)

        grand_total = response.context["grand_total"]
        self.assertEqual(grand_total["total"], 1)


    
    def test_ip_ncd_report_excludes_other_facility_records(self):
        IPNCDRecord.objects.create(
            patient=self.patient_a,
            disease="DIABETES",
            diagnosis_date=date(2026, 10, 10),
        )
        IPNCDRecord.objects.create(
            patient=self.patient_b,
            disease="DIABETES",
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("ip_ncd_report"),
            {
                "start_date": "2026-10-01",
                "end_date": "2026-10-31",
            },
        )

        self.assertEqual(response.status_code, 200)

        rows = response.context["report_rows"]
        diabetes_row = next(
            row for row in rows
            if row["name"] == "Diabetes Mellitus"
        )

        # Facility A has one male patient; Facility B must not be counted.
        self.assertEqual(diabetes_row["male"], 1)
        self.assertEqual(diabetes_row["female"], 0)
        self.assertEqual(diabetes_row["total"], 1)

        grand_total = response.context["grand_total"]
        self.assertEqual(grand_total["total"], 1)

    
    def test_anc_register_excludes_other_facility_patients(self):
        self.client.force_login(self.user_a)
        response = self.client.get(reverse("anc_register"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ANC Patient Facility A")
        self.assertNotContains(response, "ANC Patient Facility B")

    def test_cannot_view_another_facility_anc_patient(self):
        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("anc_patient_detail", args=[self.anc_patient_b.pk])
        )

        self.assertEqual(response.status_code, 404)

    def test_cannot_edit_another_facility_anc_patient(self):
        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("anc_patient_edit", args=[self.anc_patient_b.pk])
        )

        self.assertEqual(response.status_code, 404)

    def test_cannot_edit_another_facility_anc_visit(self):
        visit = ANCVisit.objects.create(
            patient=self.anc_patient_b,
            visit_date=date(2026, 10, 10),
            trimester="3RD",
        )

        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse(
                "anc_visit_edit",
                args=[self.anc_patient_b.pk, visit.pk],
            )
        )

        self.assertEqual(response.status_code, 404)

    def test_cannot_add_visit_to_another_facility_anc_patient(self):
        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("anc_visit_add", args=[self.anc_patient_b.pk])
        )

        self.assertEqual(response.status_code, 404)
    
    def test_anc_excel_report_excludes_other_facility_visits(self):
        ANCVisit.objects.create(
            patient=self.anc_patient_a,
            visit_date=date(2026, 10, 10),
            trimester="3RD",
        )
        ANCVisit.objects.create(
            patient=self.anc_patient_b,
            visit_date=date(2026, 10, 10),
            trimester="3RD",
        )

        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("anc_monthly_report_excel"),
            {"year": 2026, "month": 10},
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(
            BytesIO(response.content),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
        exported_values = [
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        workbook.close()

        self.assertIn("ANC Patient Facility A", exported_values)
        self.assertNotIn("ANC Patient Facility B", exported_values)
        self.assertIn("1", exported_values)
    
    def test_monthly_op_report_and_excel_are_facility_scoped(self):
        disease = DiseaseMaster.objects.create(
            name="Keelvayu",
            is_active=True,
        )

        DailyDiseaseCount.objects.create(
            facility=self.facility_a,
            date=date(2026, 10, 10),
            disease=disease,
            male_count=3,
            female_count=2,
            child_count=1,
        )
        DailyDiseaseCount.objects.create(
            facility=self.facility_b,
            date=date(2026, 10, 10),
            disease=disease,
            male_count=40,
            female_count=30,
            child_count=20,
        )

        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("monthly_disease_report"),
            {"year": 2026, "month": 10},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["grand_total"], 6)

        excel_response = self.client.get(
            reverse("monthly_disease_report"),
            {"year": 2026, "month": 10, "export": "excel"},
        )

        self.assertEqual(excel_response.status_code, 200)

        workbook = load_workbook(
            BytesIO(excel_response.content),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
        values = [
            cell.value
            for row in sheet.iter_rows()
            for cell in row
        ]
        workbook.close()

        self.assertIn(6, values)
        self.assertNotIn(90, values)
    
    def test_ip_fortnightly_report_excludes_other_facility_patients(self):
        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("ip_fortnightly_report"),
            {
                "start_date": "2026-10-01",
                "end_date": "2026-10-31",
            },
        )

        self.assertEqual(response.status_code, 200)

        admission_counts = response.context["admission_counts"]
        discharge_counts = response.context["discharge_counts"]

        # Facility A has one admission and no discharge in this period.
        self.assertEqual(admission_counts["total"], 1)
        self.assertEqual(discharge_counts["total"], 0)

        # Facility B's patient must not be included.
        self.assertEqual(admission_counts["male"], 1)
        self.assertEqual(admission_counts["female"], 0)
    
        excel_response = self.client.get(
            reverse("ip_fortnightly_report"),
            {
                "start_date": "2026-10-01",
                "end_date": "2026-10-31",
                "export": "excel",
            },
        )

        self.assertEqual(excel_response.status_code, 200)

        workbook = load_workbook(
            BytesIO(excel_response.content),
            read_only=True,
            data_only=True,
        )

        summary = workbook["Fortnight Summary"]
        self.assertEqual(summary["B6"].value, 1)
        self.assertEqual(summary["C6"].value, 0)

        admissions_sheet = workbook["Admissions"]
        admission_values = [
            cell.value
            for row in admissions_sheet.iter_rows()
            for cell in row
        ]
        self.assertIn("Test Patient A", admission_values)
        self.assertNotIn("Test Patient B", admission_values)

        discharges_sheet = workbook["Discharges"]
        discharge_values = [
            cell.value
            for row in discharges_sheet.iter_rows()
            for cell in row
        ]
        self.assertNotIn("Test Patient B", discharge_values)

        workbook.close()
    
    def test_daily_disease_save_is_scoped_to_assigned_facility(self):
        disease = DiseaseMaster.objects.create(
            name="Keelvayu",
            is_active=True,
        )
        selected_date = date(2026, 10, 10)

        DailyDiseaseCount.objects.create(
            facility=self.facility_b,
            date=selected_date,
            disease=disease,
            male_count=40,
            female_count=30,
            child_count=20,
        )

        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse("daily_disease_count") + "?date=2026-10-10",
            {
                f"disease_{disease.id}-male_count": "3",
                f"disease_{disease.id}-female_count": "2",
                f"disease_{disease.id}-child_count": "1",
            },
        )

        self.assertEqual(response.status_code, 302)

        record_a = DailyDiseaseCount.objects.get(
            facility=self.facility_a,
            date=selected_date,
            disease=disease,
        )
        record_b = DailyDiseaseCount.objects.get(
            facility=self.facility_b,
            date=selected_date,
            disease=disease,
        )

        self.assertEqual(record_a.male_count, 3)
        self.assertEqual(record_a.female_count, 2)
        self.assertEqual(record_a.child_count, 1)

        self.assertEqual(record_b.male_count, 40)
        self.assertEqual(record_b.female_count, 30)
        self.assertEqual(record_b.child_count, 20)
    
    def test_daily_ncd_save_is_scoped_to_assigned_facility(self):
        selected_date = date(2026, 10, 10)

        DailyNCDCount.objects.create(
            facility=self.facility_b,
            date=selected_date,
            disease="DIABETES",
            male_count=40,
            female_count=30,
            child_count=20,
        )

        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse("daily_ncd_count") + "?date=2026-10-10",
            {
                "ncd_DIABETES-male_count": "3",
                "ncd_DIABETES-female_count": "2",
                "ncd_DIABETES-child_count": "1",
            },
        )

        self.assertEqual(response.status_code, 302)

        record_a = DailyNCDCount.objects.get(
            facility=self.facility_a,
            date=selected_date,
            disease="DIABETES",
        )
        record_b = DailyNCDCount.objects.get(
            facility=self.facility_b,
            date=selected_date,
            disease="DIABETES",
        )

        self.assertEqual(record_a.male_count, 3)
        self.assertEqual(record_a.female_count, 2)
        self.assertEqual(record_a.child_count, 1)

        self.assertEqual(record_b.male_count, 40)
        self.assertEqual(record_b.female_count, 30)
        self.assertEqual(record_b.child_count, 20)
   
   



       
    

    def test_cannot_edit_another_facility_ip_disease_record(self):
        disease = DiseaseMaster.objects.create(
            name="Cross Facility Disease",
            is_active=True,
        )
        record = IPDiseaseRecord.objects.create(
            patient=self.patient_b,
            disease=disease,
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("ip_disease_edit", args=[record.pk])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(IPDiseaseRecord.objects.filter(pk=record.pk).exists())

    def test_cannot_delete_another_facility_ip_disease_record(self):
        disease = DiseaseMaster.objects.create(
            name="Cross Facility Disease Delete",
            is_active=True,
        )
        record = IPDiseaseRecord.objects.create(
            patient=self.patient_b,
            disease=disease,
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)
        response = self.client.post(
            reverse("ip_disease_delete", args=[record.pk])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(IPDiseaseRecord.objects.filter(pk=record.pk).exists())

    def test_cannot_edit_another_facility_ip_ncd_record(self):
        record = IPNCDRecord.objects.create(
            patient=self.patient_b,
            disease="DIABETES",
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)
        response = self.client.get(
            reverse("ip_ncd_edit", args=[record.pk])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(IPNCDRecord.objects.filter(pk=record.pk).exists())

    def test_cannot_delete_another_facility_ip_ncd_record(self):
        record = IPNCDRecord.objects.create(
            patient=self.patient_b,
            disease="DIABETES",
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)
        response = self.client.post(
            reverse("ip_ncd_delete", args=[record.pk])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(IPNCDRecord.objects.filter(pk=record.pk).exists())

    def test_monthly_top20_op_report_and_excel_are_facility_scoped(self):
        disease = DiseaseMaster.objects.create(
            name="Keelvayu",
            is_active=True,
        )
        DailyDiseaseCount.objects.create(
            facility=self.facility_a,
            disease=disease,
            date=date(2026, 10, 10),
            male_count=3,
            female_count=2,
            child_count=1,
        )
        DailyDiseaseCount.objects.create(
            facility=self.facility_b,
            disease=disease,
            date=date(2026, 10, 10),
            male_count=40,
            female_count=30,
            child_count=20,
        )

        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("monthly_top20_op_report"),
            {"year": 2026, "month": 10},
        )
        self.assertEqual(response.status_code, 200)
        row = next(
            r for r in response.context["report_rows"]
            if r["name"] == "Keelvayu"
        )
        self.assertEqual(row["total"], 6)
        self.assertEqual(response.context["grand_total"], 6)

        excel_response = self.client.get(
            reverse("monthly_top20_op_report"),
            {"year": 2026, "month": 10, "export": "excel"},
        )
        self.assertEqual(excel_response.status_code, 200)
        workbook = load_workbook(BytesIO(excel_response.content), data_only=True)
        sheet = workbook.active
        keelvayu_row = next(
            row for row in sheet.iter_rows(min_row=5, values_only=True)
            if row[1] == "Keelvayu"
        )
        self.assertEqual(keelvayu_row[5], 6)
        self.assertEqual(sheet.cell(sheet.max_row, 6).value, 6)

    def test_monthly_top20_ip_report_and_excel_are_facility_scoped(self):
        disease = DiseaseMaster.objects.create(
            name="Irumal",
            is_active=True,
        )
        IPDiseaseRecord.objects.create(
            patient=self.patient_a,
            disease=disease,
            diagnosis_date=date(2026, 10, 10),
        )
        IPDiseaseRecord.objects.create(
            patient=self.patient_b,
            disease=disease,
            diagnosis_date=date(2026, 10, 10),
        )

        self.client.force_login(self.user_a)

        response = self.client.get(
            reverse("monthly_top20_ip_report"),
            {"year": 2026, "month": 10},
        )
        self.assertEqual(response.status_code, 200)
        row = next(
            r for r in response.context["report_rows"]
            if r["name"] == "Irumal"
        )
        self.assertEqual(row["total"], 1)
        self.assertEqual(response.context["grand_total"], 1)

        excel_response = self.client.get(
            reverse("monthly_top20_ip_report"),
            {"year": 2026, "month": 10, "export": "excel"},
        )
        self.assertEqual(excel_response.status_code, 200)
        workbook = load_workbook(BytesIO(excel_response.content), data_only=True)
        sheet = workbook.active
        irumal_row = next(
            row for row in sheet.iter_rows(min_row=5, values_only=True)
            if row[1] == "Irumal"
        )
        self.assertEqual(irumal_row[5], 1)
        self.assertEqual(sheet.cell(sheet.max_row, 6).value, 1)
