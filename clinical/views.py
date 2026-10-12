from datetime import date
from django.contrib import messages
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .forms import (
    DailyCensusForm,
    IPPatientForm,
    IPDiseaseRecordForm,
    IPNCDRecordForm,
    DiseaseMasterForm,
    DailyDiseaseCountForm,
    DailyNCDCountForm,
    CentreProfileForm,
    ANCPatientForm,
        ANCVisitForm,
)
from .models import (
    DailyCensus,
    IPPatient,
    IPDiseaseRecord,
    IPNCDRecord,
    DiseaseMaster,
    DailyDiseaseCount,
    DailyNCDCount,
    CentreProfile,
    ANCPatient,
    ANCVisit,
     
)
from django.contrib.auth.decorators import login_required
from .facility_access import get_user_facility
from .ip_utils import calculate_ip_summary
OFFICIAL_DISEASE_ORDER = [
    "Keelvayu",
    "Irumal",
    "Soolai",
    "Karappan",
    "Gunmam",
    "Eruvaimulai",
    "Suram",
    "Peenisam",
    "Thandagavatham",
    "Kalanjagapadai",
    "Vellai",
    "Ovvamai",
    "Saganavatham",
    "Kalladaippu",
    "Neerizhivu, Neerizhivu-pun",
    "Pakkavatham",
    "Iraippu",
    "Thasaivankazhalai",
    "Maladu",
    "Perumpadu",
    "Others",
]

@login_required
def daily_census(request):
    facility = get_user_facility(request)
    today = timezone.localdate()

    selected_date = request.GET.get("date")
    if selected_date:
        try:
            selected_date = date.fromisoformat(selected_date)
        except (ValueError, TypeError):
            selected_date = today
    else:
        selected_date = today

    existing = DailyCensus.objects.filter(
        facility=facility,
        date=selected_date,
    ).first()

    if request.method == "POST":
        posted_date = request.POST.get("date")

        try:
            form_date = date.fromisoformat(posted_date)
        except (ValueError, TypeError):
            form_date = today

        existing = DailyCensus.objects.filter(
            facility=facility,
            date=form_date,
        ).first()

        if existing:
            form = DailyCensusForm(request.POST, instance=existing)
        else:
            form = DailyCensusForm(request.POST)

        if form.is_valid():
            census = form.save(commit=False)
            census.facility = facility
            census.save()

            messages.success(
                request,
                f"Daily census for {census.date.strftime('%d-%m-%Y')} "
                "saved successfully."
            )

            return redirect(f"/?date={census.date.strftime('%Y-%m-%d')}")

    else:
        if existing:
            form = DailyCensusForm(instance=existing)
        else:
            form = DailyCensusForm(initial={"date": selected_date})

    ip_summary = calculate_ip_summary(selected_date, facility)

    context = {
        "form": form,
        "census": existing,
        "selected_date": selected_date,
        "is_editing": existing is not None,
        "ip_summary": ip_summary,
        "facility": facility,
    }

    return render(request, "clinical/daily_census.html", context)



@login_required
def ip_register(request):
    facility = get_user_facility(request)

    patients = IPPatient.objects.filter(facility=facility)

    if request.method == "POST":
        form = IPPatientForm(request.POST)

        if form.is_valid():
            patient = form.save(commit=False)
            patient.facility = facility
            patient.save()

            messages.success(
                request,
                f"{patient.name} added to IP register."
            )

            return redirect("ip_register")
    else:
        form = IPPatientForm(
            initial={
                "admission_date": timezone.localdate()
            }
        )

    return render(
        request,
        "clinical/ip_register.html",
        {
            "form": form,
            "patients": patients,
            "facility": facility,
        }
    )



@login_required
def ip_disease_entry(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        form = IPDiseaseRecordForm(request.POST, facility=facility)
        if form.is_valid():
            record = form.save()
            messages.success(
                request,
                f"Disease record saved for {record.patient.name}."
            )
            return redirect("ip_disease_entry")
    else:
        form = IPDiseaseRecordForm(
            facility=facility,
            initial={"diagnosis_date": timezone.localdate()},
        )

    records = IPDiseaseRecord.objects.filter(
        patient__facility=facility
    ).select_related("patient", "disease")[:200]

    return render(
        request,
        "clinical/ip_disease_entry.html",
        {"form": form, "records": records, "facility": facility},
    )




@login_required
def ip_ncd_entry(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        form = IPNCDRecordForm(request.POST, facility=facility)
        if form.is_valid():
            record = form.save()
            messages.success(
                request,
                f"NCD record saved for {record.patient.name}."
            )
            return redirect("ip_ncd_entry")
    else:
        form = IPNCDRecordForm(
            facility=facility,
            initial={"diagnosis_date": timezone.localdate()},
        )

    records = IPNCDRecord.objects.filter(
        patient__facility=facility
    ).select_related("patient")[:200]

    return render(
        request,
        "clinical/ip_ncd_entry.html",
        {"form": form, "records": records, "facility": facility},
    )



@login_required
def ip_delete(request, pk):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        IPPatient,
        pk=pk,
        facility=facility,
    )

    if request.method == "POST":
        patient_name = patient.name
        patient.delete()
        messages.success(
            request,
            f"{patient_name} deleted from IP register."
        )
        return redirect("ip_register")

    return render(
        request,
        "clinical/ip_delete.html",
        {"patient": patient, "facility": facility},
    )
@login_required
def ip_edit(request, pk):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        IPPatient,
        pk=pk,
        facility=facility,
    )

    if request.method == "POST":
        form = IPPatientForm(request.POST, instance=patient)

        if form.is_valid():
            form.save()
            messages.success(
                request,
                f"{patient.name} updated successfully."
            )
            return redirect("ip_register")
    else:
        form = IPPatientForm(instance=patient)

    return render(
        request,
        "clinical/ip_edit.html",
        {
            "form": form,
            "patient": patient,
            "facility": facility,
        },
    )

    patient = get_object_or_404(
        IPPatient,
        pk=pk
    )

    if request.method == "POST":

        form = IPPatientForm(
            request.POST,
            instance=patient
        )

        if form.is_valid():

            patient = form.save()

            messages.success(
                request,
                f"{patient.name} updated successfully."
            )

            return redirect("ip_register")

    else:

        form = IPPatientForm(
            instance=patient
        )

    return render(
        request,
        "clinical/ip_edit.html",
        {
            "form": form,
            "patient": patient
        }
    )
@login_required
def ip_fortnightly_report(request):
    facility = get_user_facility(request)
    from datetime import timedelta
    from openpyxl.utils import get_column_letter

    today = timezone.localdate()
    default_start = today - timedelta(days=13)

    try:
        start_date = date.fromisoformat(
            request.GET.get("start_date", default_start.isoformat())
        )
        end_date = date.fromisoformat(
            request.GET.get("end_date", today.isoformat())
        )
    except (TypeError, ValueError):
        start_date = default_start
        end_date = today

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    admissions = IPPatient.objects.filter(
        facility=facility,
        admission_date__range=(start_date, end_date)
    ).order_by("admission_date", "id")

    discharges = IPPatient.objects.filter(
        facility=facility,
        discharge_date__range=(start_date, end_date)
    ).order_by("discharge_date", "id")

    def patient_counts(queryset):
        return {
            "total": queryset.count(),
            "male": queryset.filter(sex="M").count(),
            "female": queryset.filter(sex="F").count(),
            "children": queryset.filter(is_child=True).count(),
        }

    admission_counts = patient_counts(admissions)
    discharge_counts = patient_counts(discharges)

    # Historical census represents patients present, not admission events.
    register_start_date = date(2026, 10, 8)
    historical_end_date = min(
        end_date,
        register_start_date - timedelta(days=1),
    )
    historical_census = DailyCensus.objects.filter(
        facility=facility,
        date__range=(start_date, historical_end_date),
        date__lt=register_start_date,
    ).order_by("date")

    if request.GET.get("export") == "excel":
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Fortnight Summary"

        summary.append(["IP FORTNIGHTLY REPORT"])
        summary.append(["From", start_date.isoformat()])
        summary.append(["To", end_date.isoformat()])
        summary.append([])
        summary.append(["Measure", "Admissions", "Discharges"])
        summary.append([
            "Total",
            admission_counts["total"],
            discharge_counts["total"],
        ])
        summary.append([
            "Male",
            admission_counts["male"],
            discharge_counts["male"],
        ])
        summary.append([
            "Female",
            admission_counts["female"],
            discharge_counts["female"],
        ])
        summary.append([
            "Children",
            admission_counts["children"],
            discharge_counts["children"],
        ])

        for cell in summary[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="176B45"
            )

        for column in summary.columns:
            summary.column_dimensions[
                get_column_letter(column[0].column)
            ].width = 22

        for sheet_name, queryset, date_label, date_field in [
            ("Admissions", admissions, "Admission Date", "admission_date"),
            ("Discharges", discharges, "Discharge Date", "discharge_date"),
        ]:
            sheet = workbook.create_sheet(sheet_name)
            sheet.append([
                date_label, "Patient Name", "Sex", "Child",
                "Session",
            ])

            for patient in queryset:
                event_date = getattr(patient, date_field)
                session_field = (
                    "admission_session"
                    if sheet_name == "Admissions"
                    else "discharge_session"
                )
                session = getattr(patient, session_field, "") or ""
                sheet.append([
                    event_date.isoformat() if event_date else "",
                    patient.name,
                    patient.get_sex_display(),
                    "Yes" if patient.is_child else "No",
                    session,
                ])

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid", fgColor="176B45"
                )

            for column in sheet.columns:
                sheet.column_dimensions[
                    get_column_letter(column[0].column)
                ].width = 22

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

        historical_sheet = workbook.create_sheet("Historical IP Census")
        historical_sheet.append([
            "Date",
            "Morning Male",
            "Morning Female",
            "Morning Male Child",
            "Morning Female Child",
            "Morning Total",
            "Evening Male",
            "Evening Female",
            "Evening Male Child",
            "Evening Female Child",
            "Evening Total",
        ])

        for record in historical_census:
            morning_total = (
                record.ip_morning_male
                + record.ip_morning_female
                + record.ip_morning_male_child
                + record.ip_morning_female_child
            )
            evening_total = (
                record.ip_evening_male
                + record.ip_evening_female
                + record.ip_evening_male_child
                + record.ip_evening_female_child
            )
            historical_sheet.append([
                record.date.isoformat(),
                record.ip_morning_male,
                record.ip_morning_female,
                record.ip_morning_male_child,
                record.ip_morning_female_child,
                morning_total,
                record.ip_evening_male,
                record.ip_evening_female,
                record.ip_evening_male_child,
                record.ip_evening_female_child,
                evening_total,
            ])

        for cell in historical_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B45")

        for column in historical_sheet.columns:
            historical_sheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = 20

        historical_sheet.freeze_panes = "A2"
        historical_sheet.auto_filter.ref = historical_sheet.dimensions

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = (
            'attachment; filename="ip_fortnightly_report.xlsx"'
        )
        workbook.save(response)
        return response

    return render(
        request,
        "clinical/ip_fortnightly_report.html",
        {
            "start_date": start_date,
            "end_date": end_date,
            "admission_counts": admission_counts,
            "discharge_counts": discharge_counts,
            "admissions": admissions,
            "discharges": discharges,
            "historical_census": historical_census,
        },
    )


@login_required
def ip_disease_report(request):
    facility = get_user_facility(request)
    from datetime import timedelta
    from openpyxl.utils import get_column_letter

    today = timezone.localdate()
    default_start = today.replace(day=1)

    try:
        start_date = date.fromisoformat(
            request.GET.get("start_date", default_start.isoformat())
        )
        end_date = date.fromisoformat(
            request.GET.get("end_date", today.isoformat())
        )
    except (TypeError, ValueError):
        start_date = default_start
        end_date = today

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    records = IPDiseaseRecord.objects.filter(
        patient__facility=facility,
        diagnosis_date__range=(start_date, end_date)
    ).select_related("patient", "disease")

    disease_ids = records.values_list(
        "disease_id", flat=True
    ).distinct()

    diseases = DiseaseMaster.objects.filter(
        id__in=disease_ids
    ).order_by("name")

    report_rows = []
    grand_total = {
        "male": 0,
        "female": 0,
        "children": 0,
        "total": 0,
    }

    for disease in diseases:
        disease_records = records.filter(disease=disease)

        male = disease_records.filter(
            patient__sex="M",
            patient__is_child=False,
        ).values("patient_id").distinct().count()

        female = disease_records.filter(
            patient__sex="F",
            patient__is_child=False,
        ).values("patient_id").distinct().count()

        children = disease_records.filter(
            patient__is_child=True,
        ).values("patient_id").distinct().count()

        total = male + female + children

        report_rows.append({
            "disease": disease,
            "male": male,
            "female": female,
            "children": children,
            "total": total,
        })

        grand_total["male"] += male
        grand_total["female"] += female
        grand_total["children"] += children
        grand_total["total"] += total

    if request.GET.get("export") == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "IP Disease Report"

        sheet.append(["IP DISEASE CLASSIFICATION REPORT"])
        sheet.append(["From", start_date.isoformat()])
        sheet.append(["To", end_date.isoformat()])
        sheet.append([])
        sheet.append([
            "S.No.", "Disease", "Male", "Female",
            "Children", "Total"
        ])

        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="246B45"
            )
            cell.alignment = Alignment(horizontal="center")

        for index, row in enumerate(report_rows, start=1):
            sheet.append([
                index,
                row["disease"].name,
                row["male"],
                row["female"],
                row["children"],
                row["total"],
            ])

        sheet.append([
            "",
            "GRAND TOTAL",
            grand_total["male"],
            grand_total["female"],
            grand_total["children"],
            grand_total["total"],
        ])

        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

        for column in sheet.columns:
            width = max(
                len(str(cell.value or ""))
                for cell in column
            ) + 3
            sheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = min(width, 35)

        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:F{sheet.max_row - 1}"

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = (
            'attachment; filename="ip_disease_report.xlsx"'
        )
        workbook.save(response)
        return response

    return render(
        request,
        "clinical/ip_disease_report.html",
        {
            "facility": facility,
            "start_date": start_date,
            "end_date": end_date,
            "report_rows": report_rows,
            "grand_total": grand_total,
        },
    )

@login_required
def ip_ncd_report(request):
    facility = get_user_facility(request)
    from datetime import timedelta
    from openpyxl.utils import get_column_letter

    today = timezone.localdate()
    default_start = today.replace(day=1)

    try:
        start_date = date.fromisoformat(
            request.GET.get("start_date", default_start.isoformat())
        )
        end_date = date.fromisoformat(
            request.GET.get("end_date", today.isoformat())
        )
    except (TypeError, ValueError):
        start_date = default_start
        end_date = today

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    records = IPNCDRecord.objects.filter(
        patient__facility=facility,
        diagnosis_date__range=(start_date, end_date)
    ).select_related("patient")

    report_rows = []
    grand_total = {
        "male": 0,
        "female": 0,
        "children": 0,
        "total": 0,
    }

    for code, label in DailyNCDCount.NCD_CHOICES:
        ncd_records = records.filter(disease=code)

        male = ncd_records.filter(
            patient__sex="M",
            patient__is_child=False,
        ).values("patient_id").distinct().count()

        female = ncd_records.filter(
            patient__sex="F",
            patient__is_child=False,
        ).values("patient_id").distinct().count()

        children = ncd_records.filter(
            patient__is_child=True,
        ).values("patient_id").distinct().count()

        total = male + female + children

        report_rows.append({
            "name": label,
            "male": male,
            "female": female,
            "children": children,
            "total": total,
        })

        grand_total["male"] += male
        grand_total["female"] += female
        grand_total["children"] += children
        grand_total["total"] += total

    if request.GET.get("export") == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "IP NCD Report"

        sheet.append(["IP NON-COMMUNICABLE DISEASE REPORT"])
        sheet.append(["From", start_date.isoformat()])
        sheet.append(["To", end_date.isoformat()])
        sheet.append([])
        sheet.append([
            "S.No.", "NCD Category", "Male", "Female",
            "Children", "Total"
        ])

        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="246B45"
            )
            cell.alignment = Alignment(horizontal="center")

        for index, row in enumerate(report_rows, start=1):
            sheet.append([
                index,
                row["name"],
                row["male"],
                row["female"],
                row["children"],
                row["total"],
            ])

        sheet.append([
            "",
            "GRAND TOTAL",
            grand_total["male"],
            grand_total["female"],
            grand_total["children"],
            grand_total["total"],
        ])

        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

        for column in sheet.columns:
            width = max(
                len(str(cell.value or ""))
                for cell in column
            ) + 3
            sheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = min(width, 35)

        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:F{sheet.max_row - 1}"

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = (
            'attachment; filename="ip_ncd_report.xlsx"'
        )
        workbook.save(response)
        return response

    return render(
        request,
        "clinical/ip_ncd_report.html",
        {
            "facility": facility,
            "start_date": start_date,
            "end_date": end_date,
            "report_rows": report_rows,
            "grand_total": grand_total,
        },
    )

@login_required
def disease_master(request):
    diseases = DiseaseMaster.objects.all()

    if request.method == "POST":
        if not request.user.is_superuser:
            return HttpResponse(
                "Only a system administrator can modify the shared disease master.",
                status=403,
            )

        form = DiseaseMasterForm(request.POST)

        if form.is_valid():
            disease = form.save()
            messages.success(
                request,
                f"{disease.name} added successfully."
            )
            return redirect("disease_master")
    else:
        form = DiseaseMasterForm()

    return render(
        request,
        "clinical/disease_master.html",
        {
            "form": form,
            "diseases": diseases,
        },
    )
def get_disease_balance(selected_date, facility):
    op_census = DailyCensus.objects.filter(
        facility=facility,
        date=selected_date,
    ).first()

    disease_counts = DailyDiseaseCount.objects.filter(
        facility=facility,
        date=selected_date,
    )

    disease_male = sum(item.male_count for item in disease_counts)
    disease_female = sum(item.female_count for item in disease_counts)
    disease_child = sum(item.child_count for item in disease_counts)

    disease_total = disease_male + disease_female + disease_child

    if op_census:
        op_male = (
            op_census.op_new_male
            + op_census.op_old_male
        )

        op_female = (
            op_census.op_new_female
            + op_census.op_old_female
        )

        op_child = (
            op_census.op_new_male_child
            + op_census.op_new_female_child
            + op_census.op_old_male_child
            + op_census.op_old_female_child
        )

        op_total = op_male + op_female + op_child
    else:
        op_male = 0
        op_female = 0
        op_child = 0
        op_total = 0

    return {
        "op_male": op_male,
        "op_female": op_female,
        "op_child": op_child,
        "op_total": op_total,
        "disease_male": disease_male,
        "disease_female": disease_female,
        "disease_child": disease_child,
        "disease_total": disease_total,
        "male_difference": op_male - disease_male,
        "female_difference": op_female - disease_female,
        "child_difference": op_child - disease_child,
        "total_difference": op_total - disease_total,
        "is_balanced": op_total == disease_total,
    }
@login_required
def daily_disease_count(request):
    facility = get_user_facility(request)

    date_str = request.GET.get("date")
    try:
        selected_date = (
            date.fromisoformat(date_str)
            if date_str
            else timezone.localdate()
        )
    except ValueError:
        selected_date = timezone.localdate()
        messages.error(request, "Invalid date. Showing today's records.")

    disease_map = {
        disease.name: disease
        for disease in DiseaseMaster.objects.filter(
            name__in=OFFICIAL_DISEASE_ORDER,
            is_active=True,
        )
    }

    diseases = [
        disease_map[name]
        for name in OFFICIAL_DISEASE_ORDER
        if name in disease_map
    ]

    if request.method == "POST":
        for disease in diseases:
            form = DailyDiseaseCountForm(
                request.POST,
                prefix=f"disease_{disease.id}",
            )

            if form.is_valid():
                DailyDiseaseCount.objects.update_or_create(
                    facility=facility,
                    date=selected_date,
                    disease=disease,
                    defaults={
                        "male_count": form.cleaned_data["male_count"],
                        "female_count": form.cleaned_data["female_count"],
                        "child_count": form.cleaned_data["child_count"],
                    },
                )
            else:
                messages.error(
                    request,
                    f"Please check the counts entered for {disease.name}.",
                )
                return redirect(f"/diseases/daily/?date={selected_date}")

        messages.success(
            request,
            f"Disease classification saved successfully for {selected_date}.",
        )
        return redirect(f"/diseases/daily/?date={selected_date}")

    rows = []

    for disease in diseases:
        existing = DailyDiseaseCount.objects.filter(
            facility=facility,
            date=selected_date,
            disease=disease,
        ).first()

        initial = {
            "male_count": existing.male_count if existing else 0,
            "female_count": existing.female_count if existing else 0,
            "child_count": existing.child_count if existing else 0,
        }

        form = DailyDiseaseCountForm(
            initial=initial,
            prefix=f"disease_{disease.id}",
        )

        rows.append({
            "disease": disease,
            "form": form,
            "existing": existing,
        })

    balance = get_disease_balance(selected_date, facility)

    return render(
        request,
        "clinical/daily_disease_count.html",
        {
            "selected_date": selected_date,
            "rows": rows,
            "balance": balance,
            "facility": facility,
        },
    )


@login_required
def monthly_disease_report(request):
    facility = get_user_facility(request)
    from calendar import monthrange
    from openpyxl.utils import get_column_letter

    today = timezone.localdate()

    try:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))
        if not 1 <= month <= 12 or not 2000 <= year <= 2100:
            raise ValueError
    except (TypeError, ValueError):
        year, month = today.year, today.month

    start_date = date(year, month, 1)
    end_date = date(year, month, monthrange(year, month)[1])

    diseases = DiseaseMaster.objects.filter(
        name__in=OFFICIAL_DISEASE_ORDER,
        is_active=True,
    )
    disease_map = {disease.name: disease for disease in diseases}

    report_rows = []
    grand_male = grand_female = grand_child = 0

    for disease_name in OFFICIAL_DISEASE_ORDER:
        disease = disease_map.get(disease_name)
        if disease is None:
            continue

        records = DailyDiseaseCount.objects.filter(
            facility=facility,
            disease=disease,
            date__range=(start_date, end_date),
        )

        male = sum(item.male_count for item in records)
        female = sum(item.female_count for item in records)
        child = sum(item.child_count for item in records)
        total = male + female + child

        report_rows.append({
            "name": disease.name,
            "male": male,
            "female": female,
            "child": child,
            "total": total,
            "days_entered": records.filter(
                male_count__gt=0
            ).values("date").union(
                records.filter(female_count__gt=0).values("date"),
                records.filter(child_count__gt=0).values("date"),
            ).count(),
        })

        grand_male += male
        grand_female += female
        grand_child += child

    grand_total = grand_male + grand_female + grand_child

    if request.GET.get("export") == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Monthly Disease Report"
        sheet.append(["VARAM - Monthly Disease Classification Report"])
        sheet.append(["Month", start_date.strftime("%B %Y")])
        sheet.append([])
        sheet.append([
            "Sl. No.", "Disease", "Male", "Female",
            "Child", "Total", "Days Entered",
        ])

        for index, row in enumerate(report_rows, start=1):
            sheet.append([
                index, row["name"], row["male"], row["female"],
                row["child"], row["total"], row["days_entered"],
            ])

        sheet.append([
            "", "GRAND TOTAL", grand_male, grand_female,
            grand_child, grand_total, "",
        ])

        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B45")

        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

        for column in sheet.columns:
            sheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = 22

        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:G{sheet.max_row - 1}"

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = (
            'attachment; filename="monthly_disease_report.xlsx"'
        )
        workbook.save(response)
        return response

    return render(
        request,
        "clinical/monthly_disease_report.html",
        {
            "year": year,
            "month": month,
            "month_choices": [(i, __import__("calendar").month_name[i]) for i in range(1, 13)],
            "year_choices": list(range(today.year - 5, today.year + 2)),
            "month_name": start_date.strftime("%B"),
            "start_date": start_date,
            "end_date": end_date,
            "report_rows": report_rows,
            "grand_male": grand_male,
            "grand_female": grand_female,
            "grand_child": grand_child,
            "grand_total": grand_total,
        },
    )


@login_required
def daily_ncd_count(request):
    facility = get_user_facility(request)

    date_str = request.GET.get("date")
    try:
        selected_date = (
            date.fromisoformat(date_str)
            if date_str
            else timezone.localdate()
        )
    except ValueError:
        selected_date = timezone.localdate()
        messages.error(request, "Invalid date. Showing today's records.")

    ncd_choices = DailyNCDCount.NCD_CHOICES

    if request.method == "POST":
        for code, name in ncd_choices:
            form = DailyNCDCountForm(
                request.POST,
                prefix=f"ncd_{code}",
            )

            if not form.is_valid():
                messages.error(
                    request,
                    f"Please check the counts entered for {name}.",
                )
                return redirect(f"/ncd/daily/?date={selected_date}")

            DailyNCDCount.objects.update_or_create(
                facility=facility,
                date=selected_date,
                disease=code,
                defaults={
                    "male_count": form.cleaned_data["male_count"],
                    "female_count": form.cleaned_data["female_count"],
                    "child_count": form.cleaned_data["child_count"],
                },
            )

        messages.success(
            request,
            f"NCD report saved successfully for {selected_date}.",
        )
        return redirect(f"/ncd/daily/?date={selected_date}")

    rows = []

    for code, name in ncd_choices:
        existing = DailyNCDCount.objects.filter(
            facility=facility,
            date=selected_date,
            disease=code,
        ).first()

        initial = {
            "male_count": existing.male_count if existing else 0,
            "female_count": existing.female_count if existing else 0,
            "child_count": existing.child_count if existing else 0,
        }

        form = DailyNCDCountForm(
            initial=initial,
            prefix=f"ncd_{code}",
        )

        rows.append({
            "code": code,
            "name": name,
            "form": form,
            "existing": existing,
        })

    return render(
        request,
        "clinical/daily_ncd_count.html",
        {
            "selected_date": selected_date,
            "rows": rows,
            "facility": facility,
        },
    )


@login_required
def centre_profile(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        form = CentreProfileForm(request.POST, instance=facility)

        if form.is_valid():
            profile = form.save(commit=False)
            profile.pk = facility.pk
            profile.save()
            messages.success(request, "Centre profile saved successfully.")
            return redirect("centre_profile")
    else:
        form = CentreProfileForm(instance=facility)
        profile = facility

    return render(
        request,
        "clinical/centre_profile.html",
        {
            "form": form,
            "profile": profile,
            "facility": facility,
        },
    )


@login_required
def anc_register(request):
    facility = get_user_facility(request)
    patients = ANCPatient.objects.filter(facility=facility).order_by("name")

    return render(
        request,
        "clinical/anc_register.html",
        {
            "patients": patients,
            "facility": facility,
        },
    )


@login_required
def anc_patient_add(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        form = ANCPatientForm(request.POST)
        if form.is_valid():
            patient = form.save(commit=False)
            patient.facility = facility
            patient.save()
            messages.success(request, "ANC patient registered successfully.")
            return redirect("anc_register")
    else:
        form = ANCPatientForm()

    return render(
        request,
        "clinical/anc_patient_form.html",
        {"form": form, "facility": facility},
    )


@login_required
def anc_patient_edit(request, pk):
    facility = get_user_facility(request)
    patient = get_object_or_404(ANCPatient, pk=pk, facility=facility)

    if request.method == "POST":
        form = ANCPatientForm(request.POST, instance=patient)
        if form.is_valid():
            form.save()
            messages.success(request, "ANC patient details updated successfully.")
            return redirect("anc_patient_detail", pk=patient.pk)
    else:
        form = ANCPatientForm(instance=patient)

    return render(
        request,
        "clinical/anc_patient_form.html",
        {
            "form": form,
            "patient": patient,
            "editing": True,
            "facility": facility,
        },
    )


@login_required
def anc_patient_detail(request, pk):
    facility = get_user_facility(request)
    patient = get_object_or_404(ANCPatient, pk=pk, facility=facility)
    visits = patient.visits.all().order_by("-visit_date", "-id")

    return render(
        request,
        "clinical/anc_patient_detail.html",
        {
            "patient": patient,
            "visits": visits,
            "facility": facility,
        },
    )


@login_required
def anc_visit_add(request, pk):
    facility = get_user_facility(request)
    patient = get_object_or_404(ANCPatient, pk=pk, facility=facility)

    if request.method == "POST":
        form = ANCVisitForm(request.POST)
        if form.is_valid():
            visit = form.save(commit=False)
            visit.patient = patient
            visit.save()
            messages.success(request, "ANC visit saved successfully.")
            return redirect("anc_patient_detail", pk=patient.pk)
    else:
        form = ANCVisitForm(initial={"visit_date": timezone.localdate()})

    return render(
        request,
        "clinical/anc_visit_form.html",
        {"patient": patient, "form": form, "facility": facility},
    )


@login_required
def anc_visit_edit(request, pk, visit_id):
    facility = get_user_facility(request)
    patient = get_object_or_404(ANCPatient, pk=pk, facility=facility)
    visit = get_object_or_404(ANCVisit, pk=visit_id, patient=patient)

    if request.method == "POST":
        form = ANCVisitForm(request.POST, instance=visit)
        if form.is_valid():
            form.save()
            messages.success(request, "ANC visit updated successfully.")
            return redirect("anc_patient_detail", pk=patient.pk)
    else:
        form = ANCVisitForm(instance=visit)

    return render(
        request,
        "clinical/anc_visit_form.html",
        {
            "patient": patient,
            "form": form,
            "visit": visit,
            "editing": True,
            "facility": facility,
        },
    )


@login_required
def anc_monthly_report_excel(request):
    facility = get_user_facility(request)
    from datetime import date

    today = timezone.localdate()
    year = today.year
    month = today.month

    try:
        year = int(request.GET.get("year", year))
        month = int(request.GET.get("month", month))
        if month < 1 or month > 12:
            raise ValueError
        date(year, month, 1)
    except (ValueError, TypeError):
        year = today.year
        month = today.month

    month_start = date(year, month, 1)
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)

    visits = ANCVisit.objects.filter(
        patient__facility=facility,
        visit_date__gte=month_start,
        visit_date__lt=next_month,
    ).select_related("patient").order_by("visit_date", "patient__name")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ANC Monthly Report"

    sheet.append(["VARAM - ANC Monthly Report"])
    sheet.append([f"Reporting Period: {month:02d}/{year}"])
    sheet.append([])
    sheet.append([
        "Visit Date", "Beneficiary Name", "Trimester",
        "Weight (kg)", "Hb", "BP", "Kit Issued",
        "Kit Issue Date", "Kit Trimester", "Dosage",
        "Delivery Week", "Delivery Mode", "Baby Sex",
        "Baby Weight (kg)", "Neonatal Jaundice",
        "Baby Status", "Mother Status",
    ])

    header_row = 4
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="246B4A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    kit_visits = 0
    kit_trimester_counts = {
        "1ST": 0, "2ND": 0, "3RD": 0, "POSTNATAL": 0,
    }

    for visit in visits:
        sheet.append([
            visit.visit_date,
            visit.patient.name,
            visit.get_trimester_display(),
            visit.weight_kg,
            visit.hb,
            visit.bp,
            "Yes" if visit.kit_issued else "No",
            visit.kit_issue_date,
            visit.get_kit_trimester_display() or visit.kit_trimester or "",
            visit.dosage,
            visit.week_of_delivery,
            visit.mode_of_delivery,
            visit.baby_sex,
            visit.baby_weight_kg,
            "Yes" if visit.neonatal_jaundice else "No",
            visit.baby_status,
            visit.mother_status,
        ])

        if visit.kit_issued:
            kit_visits += 1
            kit_trimester = visit.kit_trimester or visit.trimester
            if kit_trimester in kit_trimester_counts:
                kit_trimester_counts[kit_trimester] += 1

    sheet.append([])
    sheet.append(["SUMMARY"])
    sheet.append(["Total ANC Visits", visits.count()])
    sheet.append([
        "Unique ANC Beneficiaries",
        visits.values("patient_id").distinct().count(),
    ])
    sheet.append(["Kit-Issued Visits", kit_visits])
    sheet.append([])
    sheet.append(["Trimester-wise Kit Distribution", "Kit-Issued Visits"])
    for key, label in [
        ("1ST", "1st Trimester"),
        ("2ND", "2nd Trimester"),
        ("3RD", "3rd Trimester"),
        ("POSTNATAL", "Post-Natal"),
    ]:
        sheet.append([label, kit_trimester_counts[key]])

    for column_cells in sheet.columns:
        max_length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12), 30
        )

    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            if hasattr(cell.value, "strftime"):
                cell.number_format = "dd-mm-yyyy"

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:Q{4 + visits.count()}"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="VARAM_ANC_Report_{year}_{month:02d}.xlsx"'
    )
    workbook.save(response)
    return response


@login_required
def anc_monthly_report(request):
    facility = get_user_facility(request)
    from datetime import date
    from .models import ANCVisit

    today = timezone.localdate()
    year = today.year
    month = today.month

    year_param = request.GET.get("year")
    month_param = request.GET.get("month")

    if year_param and month_param:
        try:
            year = int(year_param)
            month = int(month_param)
            if month < 1 or month > 12:
                raise ValueError
            date(year, month, 1)
        except ValueError:
            year = today.year
            month = today.month

    month_start = date(year, month, 1)

    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)

    visits = ANCVisit.objects.filter(
        patient__facility=facility,
        visit_date__gte=month_start,
        visit_date__lt=next_month,
    ).select_related("patient").order_by("visit_date", "patient__name")

    trimester_counts = {
        "1ST": 0,
        "2ND": 0,
        "3RD": 0,
        "POSTNATAL": 0,
    }

    kit_visits = 0
    kit_beneficiaries = set()

    kit_trimester_counts = {
        "1ST": 0,
        "2ND": 0,
        "3RD": 0,
        "POSTNATAL": 0,
    }

    for visit in visits:
        trimester_counts[visit.trimester] = (
            trimester_counts.get(visit.trimester, 0) + 1
        )

        if visit.kit_issued:
            kit_visits += 1
            kit_beneficiaries.add(visit.patient_id)

            kit_trimester = visit.kit_trimester or visit.trimester

            if kit_trimester in kit_trimester_counts:
                kit_trimester_counts[kit_trimester] += 1

    context = {
        "year": year,
        "month": month,
        "visits": visits,
        "total_visits": visits.count(),
        "unique_beneficiaries": visits.values("patient_id").distinct().count(),
        "trimester_counts": trimester_counts,
        "kit_visits": kit_visits,
        "unique_kit_beneficiaries": len(kit_beneficiaries),
        "kit_trimester_counts": kit_trimester_counts,
    }

    return render(
        request,
        "clinical/anc_monthly_report.html",
        context,
    )



@login_required
def ip_disease_edit(request, pk):
    facility = get_user_facility(request)
    record = get_object_or_404(
        IPDiseaseRecord,
        pk=pk,
        patient__facility=facility,
    )
    form = IPDiseaseRecordForm(
        request.POST or None,
        instance=record,
        facility=facility,
    )

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "IP disease record updated.")
        return redirect("ip_disease_entry")

    return render(request, "clinical/ip_record_edit.html", {
        "form": form,
        "record": record,
        "kind": "Disease",
        "back_url": "/ip/diseases/",
        "facility": facility,
    })


@login_required
def ip_ncd_edit(request, pk):
    facility = get_user_facility(request)
    record = get_object_or_404(
        IPNCDRecord,
        pk=pk,
        patient__facility=facility,
    )
    form = IPNCDRecordForm(
        request.POST or None,
        instance=record,
        facility=facility,
    )

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "IP NCD record updated.")
        return redirect("ip_ncd_entry")

    return render(request, "clinical/ip_record_edit.html", {
        "form": form,
        "record": record,
        "kind": "NCD",
        "back_url": "/ip/ncd/",
        "facility": facility,
    })


@login_required
def ip_disease_delete(request, pk):
    facility = get_user_facility(request)
    record = get_object_or_404(
        IPDiseaseRecord,
        pk=pk,
        patient__facility=facility,
    )

    if request.method == "POST":
        record.delete()
        messages.success(request, "IP disease record deleted.")
        return redirect("ip_disease_entry")

    return render(request, "clinical/ip_record_delete.html", {
        "record": record,
        "kind": "Disease",
        "description": f"{record.patient.name} - {record.disease.name}",
        "back_url": "/ip/diseases/",
        "facility": facility,
    })


@login_required
def ip_ncd_delete(request, pk):
    facility = get_user_facility(request)
    record = get_object_or_404(
        IPNCDRecord,
        pk=pk,
        patient__facility=facility,
    )

    if request.method == "POST":
        record.delete()
        messages.success(request, "IP NCD record deleted.")
        return redirect("ip_ncd_entry")

    return render(request, "clinical/ip_record_delete.html", {
        "record": record,
        "kind": "NCD",
        "description": f"{record.patient.name} - {record.get_disease_display()}",
        "back_url": "/ip/ncd/",
        "facility": facility,
    })


def _monthly_top20_context(request):
    from calendar import monthrange

    today = timezone.localdate()
    try:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))
        if not 2000 <= year <= 2100 or not 1 <= month <= 12:
            raise ValueError
    except (TypeError, ValueError):
        year, month = today.year, today.month

    start_date = date(year, month, 1)
    end_date = date(year, month, monthrange(year, month)[1])

    return {
        "year": year,
        "month": month,
        "month_name": start_date.strftime("%B"),
        "start_date": start_date,
        "end_date": end_date,
        "month_choices": [(i, __import__("calendar").month_name[i]) for i in range(1, 13)],
        "year_choices": list(range(today.year - 5, today.year + 2)),
    }


@login_required
def monthly_top20_op_report(request):
    facility = get_user_facility(request)
    from openpyxl.utils import get_column_letter

    context = _monthly_top20_context(request)
    start_date = context["start_date"]
    end_date = context["end_date"]

    disease_map = {
        item.name: item
        for item in DiseaseMaster.objects.filter(
            name__in=OFFICIAL_DISEASE_ORDER,
            is_active=True,
        )
    }

    rows = []
    grand_male = grand_female = grand_child = 0

    for disease_name in OFFICIAL_DISEASE_ORDER:
        disease = disease_map.get(disease_name)
        records = DailyDiseaseCount.objects.none()

        if disease:
            records = DailyDiseaseCount.objects.filter(
                facility=facility,
                  disease=disease,
                date__range=(start_date, end_date),
            )

        male = sum(r.male_count for r in records)
        female = sum(r.female_count for r in records)
        child = sum(r.child_count for r in records)
        rows.append({
            "name": disease_name,
            "male": male,
            "female": female,
            "child": child,
            "total": male + female + child,
        })
        grand_male += male
        grand_female += female
        grand_child += child

    context.update({
        "report_title": "Monthly Top 20 Diseases - OP",
        "report_type": "OP",
        "report_rows": rows,
        "grand_male": grand_male,
        "grand_female": grand_female,
        "grand_child": grand_child,
        "grand_total": grand_male + grand_female + grand_child,
    })

    if request.GET.get("export") == "excel":
        from openpyxl.styles import Border, Side
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Monthly OP Top 20"

        sheet.merge_cells("A1:F1")
        sheet["A1"] = "VARAM - MONTHLY TOP 20 DISEASES REPORT (OP)"
        sheet["A1"].font = Font(bold=True, size=15, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="176B45")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30

        sheet.merge_cells("A2:F2")
        sheet["A2"] = f'{context["month_name"]} - {context["year"]} | REGULAR - SIDDHA'
        sheet["A2"].font = Font(bold=True, size=12)
        sheet["A2"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A3:F3")
        sheet["A3"] = "MCH, THOOTHUKUDI"
        sheet["A3"].font = Font(bold=True, size=11)
        sheet["A3"].alignment = Alignment(horizontal="center")

        # Single header row: six aligned columns, without merged cells.
        sheet["A4"] = "Sl. No."
        sheet["B4"] = "NAME OF THE DISEASE"
        sheet["C4"] = "Male"
        sheet["D4"] = "Female"
        sheet["E4"] = "Child"
        sheet["F4"] = "Total"

        for index, row in enumerate(rows, start=1):
            sheet.append([index, row["name"], row["male"], row["female"], row["child"], row["total"]])

        sheet.append(["", "GRAND TOTAL", grand_male, grand_female, grand_child, grand_male + grand_female + grand_child])

        thin_green = Side(style="thin", color="82958A")
        for row_cells in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=6):
            for cell in row_cells:
                cell.border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)
                cell.alignment = Alignment(vertical="center", horizontal="center" if cell.column != 2 else "left")
                if cell.row == 4:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="176B45")

        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDECE2")

        widths = {"A": 10, "B": 34, "C": 14, "D": 14, "E": 14, "F": 14}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:F{sheet.max_row - 1}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:F{sheet.max_row}"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="monthly_op_top20_{context["year"]}_{context["month"]:02d}.xlsx"'
        )
        workbook.save(response)
        return response

    return render(request, "clinical/monthly_top20_report.html", context)


@login_required
def monthly_top20_ip_report(request):
    facility = get_user_facility(request)
    from openpyxl.utils import get_column_letter

    context = _monthly_top20_context(request)
    start_date = context["start_date"]
    end_date = context["end_date"]

    disease_map = {
        item.name: item
        for item in DiseaseMaster.objects.filter(
            name__in=OFFICIAL_DISEASE_ORDER,
            is_active=True,
        )
    }

    rows = []
    grand_male = grand_female = grand_child = 0

    for disease_name in OFFICIAL_DISEASE_ORDER:
        disease = disease_map.get(disease_name)
        records = IPDiseaseRecord.objects.none()

        if disease:
            records = IPDiseaseRecord.objects.filter(
                patient__facility=facility,
                  disease=disease,
                diagnosis_date__range=(start_date, end_date),
            ).select_related("patient")

        # Count each patient once per disease during the selected month.
        patient_ids = set()
        male = female = child = 0
        for record in records:
            patient = record.patient
            if patient.pk in patient_ids:
                continue
            patient_ids.add(patient.pk)

            if patient.is_child:
                child += 1
            elif patient.sex == "M":
                male += 1
            elif patient.sex == "F":
                female += 1

        rows.append({
            "name": disease_name,
            "male": male,
            "female": female,
            "child": child,
            "total": male + female + child,
        })
        grand_male += male
        grand_female += female
        grand_child += child

    context.update({
        "report_title": "Monthly Top 20 Diseases - IP",
        "report_type": "IP",
        "report_rows": rows,
        "grand_male": grand_male,
        "grand_female": grand_female,
        "grand_child": grand_child,
        "grand_total": grand_male + grand_female + grand_child,
    })

    if request.GET.get("export") == "excel":
        from openpyxl.styles import Border, Side
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Monthly IP Top 20"

        sheet.merge_cells("A1:F1")
        sheet["A1"] = "VARAM - MONTHLY TOP 20 DISEASES REPORT (IP)"
        sheet["A1"].font = Font(bold=True, size=15, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="176B45")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30

        sheet.merge_cells("A2:F2")
        sheet["A2"] = f'{context["month_name"]} - {context["year"]} | REGULAR - SIDDHA'
        sheet["A2"].font = Font(bold=True, size=12)
        sheet["A2"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A3:F3")
        sheet["A3"] = "MCH, THOOTHUKUDI"
        sheet["A3"].font = Font(bold=True, size=11)
        sheet["A3"].alignment = Alignment(horizontal="center")

        sheet.append(["Sl. No.", "NAME OF THE DISEASE", "Male", "Female", "Child", "Total"])

        for index, row in enumerate(rows, start=1):
            sheet.append([index, row["name"], row["male"], row["female"], row["child"], row["total"]])

        sheet.append(["", "GRAND TOTAL", grand_male, grand_female, grand_child, grand_male + grand_female + grand_child])

        thin_green = Side(style="thin", color="82958A")
        for row_cells in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=6):
            for cell in row_cells:
                cell.border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)
                cell.alignment = Alignment(vertical="center", horizontal="center" if cell.column != 2 else "left")
                if cell.row == 4:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="176B45")

        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDECE2")

        widths = {"A": 10, "B": 34, "C": 14, "D": 14, "E": 14, "F": 14}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:F{sheet.max_row - 1}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:F{sheet.max_row}"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="monthly_ip_top20_{context["year"]}_{context["month"]:02d}.xlsx"'
        )
        workbook.save(response)
        return response

    return render(request, "clinical/monthly_top20_report.html", context)

@login_required
def dashboard(request):
    facility = get_user_facility(request)
    return render(
        request,
        "clinical/dashboard.html",
        {"facility": facility},
    )

