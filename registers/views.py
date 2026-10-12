import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Max, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.text import slugify

from clinical.facility_access import get_user_facility
from clinical.models import DiseaseMaster
from .models import Register, RegisterEntry, RegisterField, ExternalTherapyCount, NominalPatient, OPEncounter, ExternalTherapyVisit
from .formulas import evaluate_formula


def _get_register(request, register_id):
    facility = get_user_facility(request)
    return get_object_or_404(
        Register,
        pk=register_id,
        facility=facility,
        is_active=True,
    )


@login_required
def register_list(request):
    facility = get_user_facility(request)
    registers = Register.objects.filter(
        facility=facility,
        is_active=True,
    )
    return render(request, "registers/register_list.html", {
        "registers": registers,
        "facility": facility,
    })


@login_required
def register_create(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        description = request.POST.get("description", "").strip()
        register_type = request.POST.get("register_type", "GENERAL")
        serial_prefix = request.POST.get("serial_prefix", "").strip()

        allowed_types = {choice[0] for choice in Register.REGISTER_TYPES}

        if not title:
            messages.error(request, "Register title is required.")
        elif register_type not in allowed_types:
            messages.error(request, "Select a valid register type.")
        elif Register.objects.filter(facility=facility, title=title).exists():
            messages.error(request, "A register with this title already exists.")
        else:
            register = Register.objects.create(
                facility=facility,
                title=title,
                description=description,
                register_type=register_type,
                serial_prefix=serial_prefix,
                created_by=request.user,
            )
            messages.success(request, "Register created. Add its fields next.")
            return redirect("registers:detail", register_id=register.pk)

    return render(request, "registers/register_create.html", {
        "register_types": Register.REGISTER_TYPES,
        "facility": facility,
    })


@login_required
def register_detail(request, register_id):
    register = _get_register(request, register_id)
    fields = register.fields.filter(is_active=True)
    entries = register.entries.select_related("created_by", "updated_by").all()

    return render(request, "registers/register_detail.html", {
        "register": register,
        "fields": fields,
        "entries": entries,
        "facility": register.facility,
    })


@login_required
def field_create(request, register_id):
    register = _get_register(request, register_id)

    if request.method == "POST":
        label = request.POST.get("label", "").strip()
        field_type = request.POST.get("field_type", "TEXT")
        required = request.POST.get("required") == "on"
        choices_text = request.POST.get("choices_text", "").strip()
        formula = request.POST.get("formula", "").strip()

        allowed_types = {choice[0] for choice in RegisterField.FIELD_TYPES}

        if not label:
            messages.error(request, "Field label is required.")
        elif field_type not in allowed_types:
            messages.error(request, "Select a valid field type.")
        elif field_type == "CHOICE" and not choices_text:
            messages.error(request, "Enter at least one dropdown option.")
        elif field_type == "FORMULA" and not formula:
            messages.error(request, "Enter a formula for the calculated field.")
        elif field_type == "FORMULA" and not re.fullmatch(
            r"[A-Za-z_][A-Za-z0-9_]*(?:\s*[+*/()-]\s*[A-Za-z_0-9.]+)*",
            formula,
        ):
            messages.error(request, "Formula contains unsupported characters or syntax.")
        elif field_type == "FORMULA" and any(
            name not in {
                f.field_key
                for f in register.fields.filter(is_active=True)
                if f.field_type == "NUMBER"
            }
            for name in re.findall(r"\b[A-Za-z_][A-Za-z0-9_]*\b", formula)
        ):
            messages.error(request, "Formulas may reference existing numeric fields only.")
        elif field_type == "FORMULA":
            try:
                evaluate_formula(
                    formula,
                    {
                        f.field_key: "1"
                        for f in register.fields.filter(is_active=True)
                        if f.field_type == "NUMBER"
                    },
                )
            except (ValueError, ArithmeticError):
                messages.error(request, "Please enter a valid arithmetic formula.")
            else:
                base_key = slugify(label).replace("-", "_")[:90] or "field"
                key = base_key
                suffix = 2
                while register.fields.filter(field_key=key).exists():
                    key = f"{base_key[:85]}_{suffix}"
                    suffix += 1
                position = (register.fields.aggregate(m=Max("position"))["m"] or 0) + 1
                RegisterField.objects.create(
                    register=register,
                    label=label,
                    field_key=key,
                    field_type=field_type,
                    required=False,
                    choices_text="",
                    formula=formula,
                    position=position,
                )
                messages.success(request, "Calculated field added successfully.")
                return redirect("registers:detail", register_id=register.pk)
        else:
            base_key = slugify(label).replace("-", "_")[:90] or "field"
            key = base_key
            suffix = 2
            while register.fields.filter(field_key=key).exists():
                key = f"{base_key[:85]}_{suffix}"
                suffix += 1

            position = (register.fields.aggregate(m=Max("position"))["m"] or 0) + 1

            RegisterField.objects.create(
                register=register,
                label=label,
                field_key=key,
                field_type=field_type,
                required=required,
                choices_text=choices_text,
                formula=formula,
                position=position,
            )
            messages.success(request, "Field added successfully.")
            return redirect("registers:detail", register_id=register.pk)

    return render(request, "registers/field_create.html", {
        "register": register,
        "field_types": RegisterField.FIELD_TYPES,
    })


def _validate_entry_data(register, post_data):
    values = {}
    errors = []
    fields = list(register.fields.filter(is_active=True))

    for field in fields:
        if field.field_type == "FORMULA":
            continue

        raw = post_data.get(f"field_{field.pk}", "").strip()

        if field.required and not raw:
            errors.append(f"{field.label} is required.")
            continue

        if not raw:
            values[field.field_key] = ""
            continue

        try:
            if field.field_type == "NUMBER":
                number = Decimal(raw)
                if not number.is_finite():
                    raise InvalidOperation
                values[field.field_key] = str(number)
            elif field.field_type == "DATE":
                values[field.field_key] = date.fromisoformat(raw).isoformat()
            elif field.field_type == "BOOLEAN":
                if raw not in ("true", "false"):
                    raise ValueError
                values[field.field_key] = raw == "true"
            elif field.field_type == "CHOICE":
                if raw not in field.choices_list:
                    raise ValueError
                values[field.field_key] = raw
            elif field.field_type == "TEXT":
                values[field.field_key] = raw
            else:
                errors.append(f"{field.label} uses a field type not yet supported.")
        except (ValueError, InvalidOperation, ValidationError):
            errors.append(f"Enter a valid value for {field.label}.")

    if not errors:
        for field in fields:
            if field.field_type != "FORMULA":
                continue
            try:
                result = evaluate_formula(field.formula, values)
                values[field.field_key] = str(result)
            except (ValueError, ArithmeticError, InvalidOperation) as exc:
                errors.append(f"{field.label}: {exc}")

    return values, errors


@login_required
def entry_create(request, register_id):
    register = _get_register(request, register_id)
    fields = register.fields.filter(is_active=True)

    if request.method == "POST":
        values, errors = _validate_entry_data(register, request.POST)

        raw_entry_date = request.POST.get("entry_date", "").strip()
        try:
            entry_date = date.fromisoformat(raw_entry_date) if raw_entry_date else date.today()
        except ValueError:
            entry_date = None
            errors.append("Enter a valid entry date.")

        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            with transaction.atomic():
                locked_register = Register.objects.select_for_update().get(
                    pk=register.pk,
                    facility=register.facility,
                    is_active=True,
                )
                serial = locked_register.next_serial_number

                entry = RegisterEntry.objects.create(
                    register=locked_register,
                    serial_number=serial,
                    entry_date=entry_date,
                    data=values,
                    created_by=request.user,
                    updated_by=request.user,
                )

                locked_register.next_serial_number = serial + 1
                locked_register.save(update_fields=["next_serial_number", "updated_at"])

            messages.success(request, f"Entry {entry.display_serial} saved.")
            return redirect("registers:detail", register_id=register.pk)

    return render(request, "registers/entry_form.html", {
        "register": register,
        "fields": fields,
        "entry": None,
    })


@login_required
def entry_edit(request, register_id, entry_id):
    register = _get_register(request, register_id)
    entry = get_object_or_404(
        RegisterEntry,
        pk=entry_id,
        register=register,
    )
    fields = register.fields.filter(is_active=True)

    if request.method == "POST":
        values, errors = _validate_entry_data(register, request.POST)

        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            entry.entry_date = request.POST.get("entry_date") or entry.entry_date
            entry.data = values
            entry.updated_by = request.user
            entry.save(update_fields=["entry_date", "data", "updated_by", "updated_at"])
            messages.success(request, "Entry updated.")
            return redirect("registers:detail", register_id=register.pk)

    return render(request, "registers/entry_form.html", {
        "register": register,
        "fields": fields,
        "entry": entry,
    })


@login_required
def entry_delete(request, register_id, entry_id):
    register = _get_register(request, register_id)
    entry = get_object_or_404(
        RegisterEntry,
        pk=entry_id,
        register=register,
    )

    if request.method == "POST":
        entry.delete()
        messages.success(request, "Entry deleted.")
        return redirect("registers:detail", register_id=register.pk)

    return render(request, "registers/entry_delete.html", {
        "register": register,
        "entry": entry,
    })


# External Therapy Register
@login_required
def external_therapy_daily(request):
    facility = get_user_facility(request)
    therapies = ExternalTherapyCount.THERAPY_CHOICES
    selected_date = request.GET.get("date") or request.POST.get("treatment_date") or date.today().isoformat()

    try:
        treatment_date = date.fromisoformat(selected_date)
    except (TypeError, ValueError):
        treatment_date = date.today()
        selected_date = treatment_date.isoformat()

    existing = {
        item.therapy: item
        for item in ExternalTherapyCount.objects.filter(
            facility=facility, treatment_date=treatment_date
        )
    }

    if request.method == "POST":
        errors = []
        submitted_counts = {}

        for code, label in therapies:
            raw = request.POST.get(f"count_{code}", "0").strip()
            try:
                value = int(raw or "0")
                if value < 0:
                    raise ValueError
                submitted_counts[code] = value
            except (TypeError, ValueError):
                errors.append(f"{label}: enter a whole number zero or greater.")

        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            with transaction.atomic():
                for code, label in therapies:
                    ExternalTherapyCount.objects.update_or_create(
                        facility=facility,
                        treatment_date=treatment_date,
                        therapy=code,
                        defaults={
                            "count": submitted_counts[code],
                            "created_by": request.user,
                        },
                    )
            messages.success(request, f"Therapy counts saved for {treatment_date:%d-%m-%Y}.")
            return redirect(f"{request.path}?date={treatment_date.isoformat()}")

    rows = [
        {
            "code": code,
            "label": label,
            "count": existing[code].count if code in existing else 0,
        }
        for code, label in therapies
    ]

    return render(request, "registers/external_therapy_daily.html", {
        "facility": facility,
        "treatment_date": treatment_date,
        "rows": rows,
    })


def _external_therapy_month_data(facility, month_value):
    try:
        month_start = date.fromisoformat(month_value + "-01")
    except (TypeError, ValueError):
        month_start = date.today().replace(day=1)
        month_value = month_start.strftime("%Y-%m")

    if month_start.month == 12:
        next_month = date(month_start.year + 1, 1, 1)
    else:
        next_month = date(month_start.year, month_start.month + 1, 1)

    totals = {
        item["therapy"]: item["total"] or 0
        for item in ExternalTherapyCount.objects.filter(
            facility=facility,
            treatment_date__gte=month_start,
            treatment_date__lt=next_month,
        ).values("therapy").annotate(total=Sum("count"))
    }

    rows = [
        {"code": code, "label": label, "count": totals.get(code, 0)}
        for code, label in ExternalTherapyCount.THERAPY_CHOICES
    ]
    grand_total = sum(row["count"] for row in rows)

    return month_value, rows, grand_total


@login_required
def external_therapy_report(request):
    facility = get_user_facility(request)
    month_value = request.GET.get("month") or date.today().strftime("%Y-%m")
    month_value, rows, grand_total = _external_therapy_month_data(facility, month_value)

    return render(request, "registers/external_therapy_report.html", {
        "facility": facility,
        "month_value": month_value,
        "rows": rows,
        "grand_total": grand_total,
    })


@login_required
def external_therapy_excel(request):
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    facility = get_user_facility(request)
    month_value = request.GET.get("month") or date.today().strftime("%Y-%m")
    month_value, rows, grand_total = _external_therapy_month_data(facility, month_value)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "External Therapy"
    sheet.append(["Therapy Treatment Details in Numbers"])
    sheet.append(["System", "Siddha"])
    sheet.append(["Facility", str(facility)])
    sheet.append(["Month", month_value])
    sheet.append([])
    sheet.append(["Sl. No.", "Names of External Therapy", "Total"])
    for index, row in enumerate(rows, start=1):
        sheet.append([index, row["label"], row["count"]])
    sheet.append(["", "TOTAL", grand_total])

    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet[6]:
        cell.font = Font(bold=True)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 16
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="external_therapy_{month_value}.xlsx"'
    return response


# Nominal Patient Register
@login_required
def nominal_list(request):
    facility = get_user_facility(request)
    query = request.GET.get("q", "").strip()

    patients = NominalPatient.objects.filter(facility=facility)
    if query:
        from django.db.models import Q
        patients = patients.filter(
            Q(op_number__icontains=query) | Q(name__icontains=query)
        )

    return render(request, "registers/nominal_list.html", {
        "facility": facility,
        "patients": patients[:200],
        "query": query,
    })


@login_required
def op_visit_history(request, patient_id):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        NominalPatient, pk=patient_id, facility=facility
    )
    encounters = OPEncounter.objects.filter(
        facility=facility, patient=patient
    ).prefetch_related("therapy_visits").order_by(
        "-encounter_date", "-created_at"
    )

    return render(request, "registers/op_visit_history.html", {
        "facility": facility,
        "patient": patient,
        "encounters": encounters,
    })


@login_required
def nominal_create(request):
    facility = get_user_facility(request)

    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        age_text = request.POST.get("age", "").strip()
        gender = request.POST.get("gender", "").strip()
        registered_date_text = request.POST.get("registered_date", "").strip()

        errors = []
        if not name:
            errors.append("Patient name is required.")

        try:
            age = int(age_text)
            if not 0 <= age <= 120:
                raise ValueError
        except (TypeError, ValueError):
            age = None
            errors.append("Enter a valid age from 0 to 120.")

        if gender not in ("M", "F"):
            errors.append("Select Male or Female.")

        try:
            registered_date = (
                date.fromisoformat(registered_date_text)
                if registered_date_text else date.today()
            )
        except ValueError:
            registered_date = date.today()
            errors.append("Enter a valid registration date.")

        if errors:
            for error in errors:
                messages.error(request, error)
            return render(request, "registers/nominal_form.html", {
                "facility": facility,
                "form_data": request.POST,
                "today": date.today(),
            })

        with transaction.atomic():
            last_serial = NominalPatient.objects.filter(
                facility=facility
            ).aggregate(max_serial=Max("serial_number"))["max_serial"] or 0
            serial_number = last_serial + 1

            prefix = facility.facility_code.strip() if facility.facility_code else ""
            op_number = f"{prefix + '-' if prefix else ''}OP-{serial_number:05d}"

            patient = NominalPatient.objects.create(
                facility=facility,
                serial_number=serial_number,
                op_number=op_number,
                name=name,
                age=age,
                gender=gender,
                diagnosis="",
                medicine="",
                registered_date=registered_date,
                created_by=request.user,
            )

        messages.success(
            request,
            f"Patient registered. S.No. {patient.serial_number}, OP No. {patient.op_number}. "
            f"{'Child' if patient.is_child else 'Adult'} category assigned automatically."
        )
        return redirect("registers:op_encounter_create", patient_id=patient.pk)

    return render(request, "registers/nominal_form.html", {
        "facility": facility,
        "form_data": {},
        "today": date.today(),
    })


@login_required
def op_encounter_create(request, patient_id):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        NominalPatient, pk=patient_id, facility=facility
    )
    diagnoses = DiseaseMaster.objects.filter(is_active=True).order_by("name")

    if request.method == "POST":
        encounter_date_text = request.POST.get("encounter_date", "").strip()
        session = request.POST.get("session", "").strip()
        case_type = request.POST.get("case_type", "").strip()
        diagnosis_id = request.POST.get("diagnosis", "").strip()
        medicine = request.POST.get("medicine", "").strip()

        try:
            encounter_date = date.fromisoformat(encounter_date_text)
        except ValueError:
            encounter_date = None

        diagnosis = diagnoses.filter(pk=diagnosis_id).first() if diagnosis_id.isdigit() else None

        if not encounter_date or session not in ("MORNING", "EVENING") or case_type not in ("NEW", "OLD"):
            messages.error(request, "Enter a valid date, session and case type.")
        elif not diagnosis:
            messages.error(request, "Select a valid diagnosis.")
        elif case_type == "NEW" and patient.registered_date != encounter_date:
            messages.error(request, "A new case must use the patient's registration date. Select Old case for a later visit.")
        else:
            encounter, created = OPEncounter.objects.get_or_create(
                facility=facility,
                patient=patient,
                encounter_date=encounter_date,
                session=session,
                defaults={
                    "case_type": case_type,
                    "diagnosis": diagnosis.name,
                    "medicine": medicine,
                    "created_by": request.user,
                },
            )
            if not created:
                warning = (
                    "An OP encounter already exists for this date and session. "
                    "Its saved details were not overwritten. Use Visit History "
                    "to edit the existing encounter."
                )
                return render(request, "registers/op_encounter_form.html", {
                    "facility": facility,
                    "patient": patient,
                    "diagnoses": diagnoses,
                    "today": date.today(),
                    "encounter_warning": warning,
                })

            messages.success(request, "OP encounter, diagnosis and medicine saved.")
            return redirect("registers:patient_therapy_create", patient_id=patient.pk)

    return render(request, "registers/op_encounter_form.html", {
        "facility": facility,
        "patient": patient,
        "diagnoses": diagnoses,
        "today": date.today(),
    })


@login_required
def op_encounter_edit(request, patient_id, encounter_id):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        NominalPatient, pk=patient_id, facility=facility
    )
    encounter = get_object_or_404(
        OPEncounter,
        pk=encounter_id,
        patient=patient,
        facility=facility,
    )
    diagnoses = DiseaseMaster.objects.filter(is_active=True).order_by("name")

    if request.method == "POST":
        diagnosis_id = request.POST.get("diagnosis", "").strip()
        medicine = request.POST.get("medicine", "").strip()
        diagnosis = (
            diagnoses.filter(pk=diagnosis_id).first()
            if diagnosis_id.isdigit()
            else None
        )

        if not diagnosis:
            messages.error(request, "Select a valid diagnosis.")
        else:
            encounter.diagnosis = diagnosis.name
            encounter.medicine = medicine
            encounter.save(update_fields=["diagnosis", "medicine"])
            messages.success(request, "OP encounter diagnosis and medicine updated.")
            return redirect("registers:op_encounter_create", patient_id=patient.pk)

    selected_diagnosis = diagnoses.filter(name=encounter.diagnosis).first()

    return render(request, "registers/op_encounter_edit.html", {
        "facility": facility,
        "patient": patient,
        "encounter": encounter,
        "diagnoses": diagnoses,
        "selected_diagnosis": selected_diagnosis,
    })


@login_required
def patient_therapy_create(request, patient_id):
    facility = get_user_facility(request)
    patient = get_object_or_404(
        NominalPatient, pk=patient_id, facility=facility
    )
    encounters = OPEncounter.objects.filter(
        facility=facility, patient=patient
    ).order_by("-encounter_date", "-created_at")

    if request.method == "POST":
        encounter_id = request.POST.get("encounter_id", "").strip()
        therapy_codes = request.POST.getlist("therapies")
        treatment_date_text = request.POST.get("treatment_date", "").strip()

        encounter = get_object_or_404(
            OPEncounter, pk=encounter_id, facility=facility, patient=patient
        ) if encounter_id.isdigit() else None

        try:
            treatment_date = date.fromisoformat(treatment_date_text)
        except ValueError:
            treatment_date = None

        allowed_codes = {code for code, label in ExternalTherapyCount.THERAPY_CHOICES}
        if not encounter:
            messages.error(request, "Select a valid OP encounter.")
        elif not treatment_date:
            messages.error(request, "Enter a valid treatment date.")
        elif not therapy_codes or any(code not in allowed_codes for code in therapy_codes):
            messages.error(request, "Select at least one valid therapy.")
        elif treatment_date != encounter.encounter_date:
            messages.error(request, "The treatment date must match the selected OP encounter date.")
        else:
            with transaction.atomic():
                for code in set(therapy_codes):
                    ExternalTherapyVisit.objects.get_or_create(
                        facility=facility,
                        encounter=encounter,
                        therapy=code,
                        treatment_date=treatment_date,
                        defaults={"created_by": request.user},
                    )
            messages.success(request, "Patient therapy selections saved.")
            return redirect("registers:nominal_list")

    visits = ExternalTherapyVisit.objects.filter(
        facility=facility, encounter__patient=patient
    ).select_related("encounter").order_by("-treatment_date")

    return render(request, "registers/patient_therapy_form.html", {
        "facility": facility,
        "patient": patient,
        "encounters": encounters,
        "visits": visits,
        "therapies": ExternalTherapyCount.THERAPY_CHOICES,
        "today": date.today(),
    })
