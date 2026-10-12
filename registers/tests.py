from datetime import date

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from clinical.models import CentreProfile, FacilityMembership
from registers.models import Register, RegisterEntry, RegisterField


User = get_user_model()


class RegisterBuilderTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.facility_a = CentreProfile.objects.create(
            centre_name="Test Hospital A",
            district="Thoothukudi",
        )
        cls.facility_b = CentreProfile.objects.create(
            centre_name="Test Hospital B",
            district="Tirunelveli",
        )

        cls.user_a = User.objects.create_user(
            username="register_user_a",
            password="TestPass123!",
        )
        cls.user_b = User.objects.create_user(
            username="register_user_b",
            password="TestPass123!",
        )

        FacilityMembership.objects.create(
            user=cls.user_a,
            facility=cls.facility_a,
            role="FACILITY_ADMIN",
        )
        FacilityMembership.objects.create(
            user=cls.user_b,
            facility=cls.facility_b,
            role="FACILITY_ADMIN",
        )

        cls.register_a = Register.objects.create(
            facility=cls.facility_a,
            title="Test OP Register",
            register_type="OP",
            serial_prefix="OP-",
            created_by=cls.user_a,
        )
        cls.field_a = RegisterField.objects.create(
            register=cls.register_a,
            label="Patient Name",
            field_key="patient_name",
            field_type="TEXT",
            required=True,
        )

    def test_register_list_shows_only_assigned_facility(self):
        self.client.force_login(self.user_b)

        response = self.client.get(reverse("registers:list"))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Test OP Register")

    def test_other_facility_register_is_not_accessible(self):
        self.client.force_login(self.user_b)

        response = self.client.get(
            reverse("registers:detail", args=[self.register_a.pk])
        )

        self.assertEqual(response.status_code, 404)

    def test_user_can_create_register_for_own_facility(self):
        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse("registers:create"),
            {
                "title": "Finance Ledger",
                "description": "Daily expenses",
                "register_type": "FINANCE",
                "serial_prefix": "FIN-",
            },
        )

        self.assertEqual(response.status_code, 302)
        register = Register.objects.get(
            facility=self.facility_a,
            title="Finance Ledger",
        )
        self.assertEqual(register.created_by, self.user_a)
        self.assertEqual(register.serial_prefix, "FIN-")

    def test_entry_serial_numbers_increment(self):
        self.client.force_login(self.user_a)
        url = reverse(
            "registers:entry_create",
            args=[self.register_a.pk],
        )

        first = self.client.post(
            url,
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "Patient One",
            },
        )
        second = self.client.post(
            url,
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "Patient Two",
            },
        )

        self.assertEqual(first.status_code, 302)
        self.assertEqual(second.status_code, 302)

        entries = RegisterEntry.objects.filter(
            register=self.register_a
        ).order_by("serial_number")

        self.assertEqual(entries.count(), 2)
        self.assertEqual(
            list(entries.values_list("serial_number", flat=True)),
            [1, 2],
        )
        self.assertEqual(entries.first().display_serial, "OP-00001")
        self.assertEqual(entries.last().display_serial, "OP-00002")

    def test_other_facility_cannot_edit_or_delete_entry(self):
        entry = RegisterEntry.objects.create(
            register=self.register_a,
            serial_number=1,
            entry_date=date(2026, 10, 11),
            data={"patient_name": "Protected Patient"},
            created_by=self.user_a,
            updated_by=self.user_a,
        )
        self.client.force_login(self.user_b)

        edit_response = self.client.get(
            reverse(
                "registers:entry_edit",
                args=[self.register_a.pk, entry.pk],
            )
        )
        delete_response = self.client.post(
            reverse(
                "registers:entry_delete",
                args=[self.register_a.pk, entry.pk],
            )
        )

        self.assertEqual(edit_response.status_code, 404)
        self.assertEqual(delete_response.status_code, 404)
        self.assertTrue(RegisterEntry.objects.filter(pk=entry.pk).exists())

    def test_required_field_rejects_empty_entry(self):
        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse(
                "registers:entry_create",
                args=[self.register_a.pk],
            ),
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            RegisterEntry.objects.filter(register=self.register_a).count(),
            0,
        )

    def test_invalid_entry_date_is_rejected(self):
        self.client.force_login(self.user_a)

        response = self.client.post(
            reverse(
                "registers:entry_create",
                args=[self.register_a.pk],
            ),
            {
                "entry_date": "not-a-date",
                f"field_{self.field_a.pk}": "Patient Three",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a valid entry date.")
        self.assertEqual(
            RegisterEntry.objects.filter(register=self.register_a).count(),
            0,
        )
from registers.formulas import evaluate_formula


class RegisterFormulaTests(TestCase):
    def test_formula_calculates_invoice_amount(self):
        result = evaluate_formula(
            "quantity * rate",
            {"quantity": "3", "rate": "25"},
        )
        self.assertEqual(str(result), "75")

    def test_formula_calculates_closing_stock(self):
        result = evaluate_formula(
            "opening + received - issued",
            {"opening": "10", "received": "5", "issued": "3"},
        )
        self.assertEqual(str(result), "12")

    def test_formula_rejects_division_by_zero(self):
        with self.assertRaises(ValueError):
            evaluate_formula("quantity / 0", {"quantity": "5"})

    def test_formula_rejects_unsafe_python(self):
        with self.assertRaises(ValueError):
            evaluate_formula(
                '__import__("os").system("whoami")',
                {},
            )
class RegisterFormulaWorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.facility_a = CentreProfile.objects.create(
            centre_name="Formula Test Hospital",
            district="Thoothukudi",
        )
        cls.user_a = User.objects.create_user(
            username="formula_user",
            password="TestPass123!",
        )
        FacilityMembership.objects.create(
            user=cls.user_a,
            facility=cls.facility_a,
            role="FACILITY_ADMIN",
        )
        cls.register_a = Register.objects.create(
            facility=cls.facility_a,
            title="Formula Test Register",
            register_type="GENERAL",
            created_by=cls.user_a,
        )
        cls.field_a = RegisterField.objects.create(
            register=cls.register_a,
            label="Patient Name",
            field_key="patient_name",
            field_type="TEXT",
            required=True,
        )

    def setUp(self):
        self.client.force_login(self.user_a)

        self.quantity_field = RegisterField.objects.create(
            register=self.register_a,
            label="Quantity",
            field_key="quantity",
            field_type="NUMBER",
            required=True,
        )
        self.rate_field = RegisterField.objects.create(
            register=self.register_a,
            label="Rate",
            field_key="rate",
            field_type="NUMBER",
            required=True,
        )
        self.amount_field = RegisterField.objects.create(
            register=self.register_a,
            label="Amount",
            field_key="amount",
            field_type="FORMULA",
            formula="quantity * rate",
            required=False,
        )

    def test_entry_saves_calculated_amount(self):
        response = self.client.post(
            reverse("registers:entry_create", args=[self.register_a.pk]),
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "Patient One",
                f"field_{self.quantity_field.pk}": "3",
                f"field_{self.rate_field.pk}": "25",
            },
        )

        self.assertEqual(response.status_code, 302)
        entry = RegisterEntry.objects.get(register=self.register_a)
        self.assertEqual(entry.data["quantity"], "3")
        self.assertEqual(entry.data["rate"], "25")
        self.assertEqual(entry.data["amount"], "75")

    def test_entry_rejects_formula_division_by_zero(self):
        self.amount_field.formula = "quantity / rate"
        self.amount_field.save(update_fields=["formula"])

        response = self.client.post(
            reverse("registers:entry_create", args=[self.register_a.pk]),
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "Patient Two",
                f"field_{self.quantity_field.pk}": "5",
                f"field_{self.rate_field.pk}": "0",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            RegisterEntry.objects.filter(register=self.register_a).count(),
            0,
        )

    def test_submitted_amount_cannot_override_formula(self):
        response = self.client.post(
            reverse("registers:entry_create", args=[self.register_a.pk]),
            {
                "entry_date": "2026-10-11",
                f"field_{self.field_a.pk}": "Patient Three",
                f"field_{self.quantity_field.pk}": "3",
                f"field_{self.rate_field.pk}": "25",
                f"field_{self.amount_field.pk}": "999999",
            },
        )

        self.assertEqual(response.status_code, 302)
        entry = RegisterEntry.objects.get(register=self.register_a)
        self.assertEqual(entry.data["amount"], "75")
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from clinical.models import CentreProfile, FacilityMembership
from registers.models import ExternalTherapyCount


class ExternalTherapyRegisterTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.user = User.objects.create_user(
            username="therapy_user", password="testpass123"
        )
        cls.other_user = User.objects.create_user(
            username="other_therapy_user", password="testpass123"
        )

        cls.facility = CentreProfile.objects.create(
            centre_name="Therapy Test Facility"
        )
        cls.other_facility = CentreProfile.objects.create(
            centre_name="Other Therapy Facility"
        )

        FacilityMembership.objects.create(
            user=cls.user, facility=cls.facility, is_active=True
        )
        FacilityMembership.objects.create(
            user=cls.other_user, facility=cls.other_facility, is_active=True
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_daily_page_loads(self):
        response = self.client.get(
            reverse("registers:external_therapy_daily")
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["rows"]), 12)

    def test_daily_counts_save_and_update(self):
        url = reverse("registers:external_therapy_daily")
        data = {
            "treatment_date": "2026-10-10",
            "count_KATTUTHAL": "2",
            "count_PATRU": "6",
            "count_OTTRADAM": "5",
            "count_PURA_VALAYAM": "3",
            "count_VARAM": "2",
            "count_POOCHU": "12",
            "count_VEDHU": "2",
            "count_THOKKANAM": "5",
            "count_PUGAI": "26",
            "count_KALIMBU": "2",
            "count_SUTTIKAI": "3",
            "count_OTHERS": "3",
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            ExternalTherapyCount.objects.filter(
                facility=self.facility,
                treatment_date=date(2026, 10, 10),
            ).count(),
            12,
        )

        data["count_PUGAI"] = "30"
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            ExternalTherapyCount.objects.get(
                facility=self.facility,
                treatment_date=date(2026, 10, 10),
                therapy="PUGAI",
            ).count,
            30,
        )
        self.assertEqual(
            ExternalTherapyCount.objects.filter(
                facility=self.facility,
                treatment_date=date(2026, 10, 10),
            ).count(),
            12,
        )

    def test_monthly_report_shows_all_therapies_and_total(self):
        ExternalTherapyCount.objects.create(
            facility=self.facility,
            therapy="KATTUTHAL",
            treatment_date=date(2026, 10, 1),
            count=2,
            created_by=self.user,
        )
        ExternalTherapyCount.objects.create(
            facility=self.facility,
            therapy="PATRU",
            treatment_date=date(2026, 10, 2),
            count=6,
            created_by=self.user,
        )
        response = self.client.get(
            reverse("registers:external_therapy_report"),
            {"month": "2026-10"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["rows"]), 12)
        self.assertEqual(response.context["grand_total"], 8)
        self.assertEqual(response.context["rows"][0]["count"], 2)
        self.assertEqual(response.context["rows"][2]["count"], 0)

    def test_monthly_report_is_facility_isolated(self):
        ExternalTherapyCount.objects.create(
            facility=self.other_facility,
            therapy="KATTUTHAL",
            treatment_date=date(2026, 10, 1),
            count=99,
            created_by=self.other_user,
        )
        response = self.client.get(
            reverse("registers:external_therapy_report"),
            {"month": "2026-10"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["grand_total"], 0)

    def test_excel_export_contains_all_therapies_and_total(self):
        ExternalTherapyCount.objects.create(
            facility=self.facility,
            therapy="KATTUTHAL",
            treatment_date=date(2026, 10, 1),
            count=4,
            created_by=self.user,
        )
        response = self.client.get(
            reverse("registers:external_therapy_excel"),
            {"month": "2026-10"},
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook["External Therapy"]
        self.assertEqual(sheet.max_row, 19)
        self.assertEqual(sheet["B7"].value, "Kattuthal (Bandage Dressing)")
        self.assertEqual(sheet["C7"].value, 4)
        self.assertEqual(sheet["B19"].value, "TOTAL")
        self.assertEqual(sheet["C19"].value, 4)
        workbook.close()

